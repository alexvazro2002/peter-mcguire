import base64
import json
import stat
import threading
from pathlib import Path

import pytest

from app.assistant.mail_review import EmailReview, EmailReviewError, participant_match


def email(message_id='a1', sender='José Angel <jose.angel@example.test>', cc='', **kwargs):
    return dict(id=message_id, thread_id=message_id, sender=sender, to='Me <me@example.test>', cc=cc,
                subject='Review ' + message_id, date='Thu, 15 Oct 2026 12:00:00 +0100', received_ms=1,
                url='https://mail.google.com/mail/u/?authuser=me@example.test#all/' + message_id,
                body='Please review the budget proposal for Project Lantern. The total is £215.',
                attachments=[], attachment_inventory_complete=True, **kwargs)


class Gmail:
    def __init__(self, messages=None):
        self.messages = messages if messages is not None else [email()]
        self.reads, self.downloads, self.lists = [], [], []
        self.account, self.complete = 'me@example.test', True
        self.attachment_bytes = b'The repayment deadline is 24 October.'

    def status(self):
        return dict(account=self.account, connected=True)

    def list_headers(self, query, **kwargs):
        self.lists.append((query, kwargs))
        return dict(account=self.account, messages=[{k: v for k, v in m.items() if k not in {'body', 'attachments'}}
                                                  for m in self.messages], complete=self.complete)

    def read(self, message_id, **kwargs):
        self.reads.append(message_id)
        return dict(account=self.account, message=next(m for m in self.messages if m['id'] == message_id))

    def read_attachment(self, message_id, attachment_id, **kwargs):
        self.downloads.append((message_id, attachment_id, kwargs))
        return self.attachment_bytes


@pytest.mark.parametrize('sender,cc', [
    ('José Angel <jose.angel@example.test>', ''),
    ('JoseAngel <joseangel@example.test>', ''),
    ('Joseangell <joseangell@example.test>', ''),
    ('Someone Else <else@example.test>', 'José Ángel <joseangel@example.test>'),
    ('Someone Else <else@example.test>', ['Other <other@example.test>', 'Joseangel@example.test']),
])
def test_participant_matches_collapsed_accented_and_near_names_in_from_or_cc(sender, cc):
    match = participant_match('Jose Angel', dict(sender=sender, cc=cc))
    assert match and match['score'] >= .86
    assert match['field'] == ('cc' if cc else 'sender')


def test_participant_does_not_match_body_or_to_or_one_shared_name():
    assert participant_match('Jose Angel', dict(sender='Jose Smith <jose@example.test>',
                                               to='Jose Angel <joseangel@example.test>',
                                               body='Jose Angel', cc='')) is None


def test_fuzzy_display_name_works_when_job_title_and_surname_follow_it():
    assert participant_match('Josie Angel',dict(sender='"José Ángel Vazquez. Director H10 Palace" <office@example.test>',cc=''))
    assert participant_match('Jose Angel',dict(sender='"Josie Smith. Hotel Manager" <office@example.test>',cc='')) is None


def test_prepare_enumerates_complete_window_then_locally_matches_from_or_cc(tmp_path):
    gmail = Gmail([email(), email('a2', sender='Else <else@example.test>', cc='JoseAngel <ja@example.test>'),
                   email('a3', sender='Different Person <different@example.test>')])
    review = EmailReview(tmp_path, gmail)
    result = review.prepare('after:1791932400 before:1792105200 -in:sent', participant='Jose Angel',
                            date_label='the past two weeks')
    assert gmail.lists[0][1]['limit'] is None
    assert gmail.reads == ['a1', 'a2']
    assert result['manifest']['headers_scanned'] == 3
    assert result['manifest']['matched_total'] == 2
    assert result['messages'][1]['match']['field'] == 'cc'
    assert len(result['sources']) == 2
    assert 'past two weeks' in result['answer']
    assert Path(result['text_directory']).is_dir()
    retrieved = review.search('Project Lantern total')
    assert '£215' in retrieved['context']
    assert {r['message_id'] for r in retrieved['passages']} == {'a1', 'a2'}
    assert retrieved['sources'][0]['id'] == 'a1'
    assert review.read('a1')['body'].startswith('Please review')
    assert not (tmp_path / 'notes').exists()


def test_preparation_rejects_unbounded_or_incomplete_enumeration(tmp_path):
    gmail = Gmail()
    review = EmailReview(tmp_path, gmail)
    with pytest.raises(EmailReviewError, match='bounded date range'):
        review.prepare('from:Jose')
    gmail.complete = False
    with pytest.raises(EmailReviewError, match='complete email date window'):
        review.prepare('after:1791932400')
    assert gmail.reads == []


def test_readable_attachment_is_saved_safely_indexed_and_private(tmp_path):
    message = email()
    message['attachments'] = [dict(filename='../../sensitive.txt', mime_type='text/plain', size=35, attachment_id='att1')]
    gmail = Gmail([message])
    review = EmailReview(tmp_path, gmail)
    result = review.prepare('after:1791932400')
    attachment = result['messages'][0]['attachments'][0]
    assert attachment['status'] == 'readable'
    directory = Path(result['cache_directory'])
    assert (directory / attachment['file']).is_file()
    assert not (tmp_path / 'sensitive.txt').exists()
    assert review.search('repayment deadline')['passages'][0]['kind'] == 'attachment: ../../sensitive.txt'
    assert '24 October' in review.search('repayment deadline')['context']
    assert result['attachment_readable'] == 1
    assert gmail.downloads[0][2]['max_bytes'] == 20 * 1024 * 1024
    for path in directory.rglob('*'):
        assert stat.S_IMODE(path.stat().st_mode) == (0o700 if path.is_dir() else 0o600)


def test_inline_attachment_and_truncated_inline_inventory_are_truthful(tmp_path):
    message = email()
    message['attachments'] = [dict(filename='list.txt', mime_type='text/plain', size=13,
                                   inline_data=base64.urlsafe_b64encode(b'Check invoice').decode()),
                              dict(filename='partial.txt', mime_type='text/plain', size=13,
                                   inline_data='aGVsbG8', inline_data_truncated=True)]
    result = EmailReview(tmp_path, Gmail([message])).prepare('after:1791932400')
    assert [a['status'] for a in result['messages'][0]['attachments']] == ['readable', 'unavailable']
    assert 'open the relevant email' in result['answer']
    assert 'no attachments' not in result['answer']


@pytest.mark.parametrize('complete', [False, True])
def test_no_attachments_only_confirmed_after_complete_full_mime_inventory(tmp_path, complete):
    message = email()
    message['attachment_inventory_complete'] = complete
    result = EmailReview(tmp_path, Gmail([message])).prepare('after:1791932400')
    assert ('Gmail reports no attachments' in result['answer']) == complete
    assert ('open the relevant email' in result['answer']) != complete


def test_unsupported_large_and_unreadable_attachments_are_not_claimed_readable(tmp_path):
    message = email()
    message['attachments'] = [dict(filename='unknown.exe', size=4, attachment_id='exec'),
                              dict(filename='huge.txt', size=9999, attachment_id='huge')]
    gmail = Gmail([message])
    result = EmailReview(tmp_path, gmail, max_attachment_bytes=100).prepare('after:1791932400')
    assert [a['status'] for a in result['messages'][0]['attachments']] == ['unsupported', 'too_large']
    assert len(gmail.downloads) == 1
    assert result['attachment_readable'] == 0
    assert 'could not be checked or read' in result['answer']


def test_expiration_removes_private_corpus_and_reloading_preserves_usable_index(tmp_path):
    now = [1000]
    gmail = Gmail()
    review = EmailReview(tmp_path, gmail, clock=lambda: now[0], ttl_seconds=60)
    result = review.prepare('after:1')
    second = EmailReview(tmp_path, gmail, clock=lambda: now[0], ttl_seconds=60)
    assert second.load()['corpus_id'] == result['corpus_id']
    assert '£215' in second.search('budget total')['context']
    now[0] += 61
    assert second.load() is None
    assert not Path(result['cache_directory']).exists()
    with pytest.raises(EmailReviewError, match='No active email review'):
        second.search('budget')


def test_account_change_cannot_reuse_another_account_cached_mail(tmp_path):
    gmail = Gmail()
    review = EmailReview(tmp_path, gmail)
    review.prepare('after:1')
    gmail.account = 'other@example.test'
    assert EmailReview(tmp_path, gmail).load() is None
    with pytest.raises(EmailReviewError, match='No active email review'):
        review.search('budget')


def test_selecting_verified_header_does_not_need_a_new_search(tmp_path):
    gmail = Gmail()
    review = EmailReview(tmp_path, gmail)
    result = review.prepare_messages([email()], account=gmail.account)
    assert gmail.lists == []
    assert result['manifest']['matched_total'] == 1
    assert review.read('a1')['body_status'] == 'readable'
    with pytest.raises(EmailReviewError, match='not part'):
        review.read('unverified')


def test_bounded_download_is_reported_as_partial_not_all_messages(tmp_path):
    gmail = Gmail([email('a1'), email('a2')])
    result = EmailReview(tmp_path, gmail, max_messages=1).prepare('after:1')
    assert result['manifest']['matched_total'] == 2
    assert result['evidence']['complete'] is False
    assert 'matched 2' in result['answer']
    assert gmail.reads == ['a1']


def test_failed_body_download_keeps_verified_subject_but_never_claims_body_read(tmp_path):
    gmail = Gmail()
    gmail.read = lambda *args, **kwargs: (_ for _ in ()).throw(OSError('failure'))
    result = EmailReview(tmp_path, gmail).prepare('after:1')
    assert result['messages'][0]['body_status'] == 'failed'
    assert result['sources'][0]['subject'] == 'Review a1'
    assert 'could not be downloaded' in result['answer']
    assert 'no attachments' not in result['answer']


def test_body_from_another_account_is_not_cached(tmp_path):
    gmail = Gmail()
    gmail.read = lambda *args, **kwargs: dict(account='other@example.test', message=email())
    result = EmailReview(tmp_path, gmail).prepare('after:1')
    assert result['messages'][0]['body_status'] == 'failed'
    assert not list(Path(result['text_directory']).iterdir())


def test_cancelled_preparation_leaves_no_partial_email_corpus(tmp_path):
    cancel = threading.Event()
    gmail = Gmail([email('a1'), email('a2')])
    original = gmail.read
    def read(*args, **kwargs):
        result = original(*args, **kwargs)
        cancel.set()
        return result
    gmail.read = read
    review = EmailReview(tmp_path, gmail)
    with pytest.raises(EmailReviewError, match='cancelled'):
        review.prepare('after:1', cancel=cancel)
    assert not list(review.base.glob('review-*'))


def test_cache_does_not_follow_symlinked_runtime_directory(tmp_path):
    outside = tmp_path / 'outside'
    outside.mkdir()
    (tmp_path / '.runtime').symlink_to(outside, target_is_directory=True)
    with pytest.raises(EmailReviewError, match='symbolic links'):
        EmailReview(tmp_path, Gmail()).prepare('after:1')
    assert not list(outside.iterdir())


def test_malformed_or_inconsistent_inline_attachment_is_never_read(tmp_path):
    message = email()
    message['attachments'] = [dict(filename='bad.txt', inline_data='%%%broken%%%', size=10),
                              dict(filename='wrong-size.txt', inline_data='aGVsbG8', size=20)]
    result = EmailReview(tmp_path, Gmail([message])).prepare('after:1')
    assert [a['status'] for a in result['messages'][0]['attachments']] == ['failed', 'failed']
    assert result['attachment_readable'] == 0


def test_account_switch_during_attachment_read_discards_the_preparation(tmp_path):
    message = email()
    message['attachments'] = [dict(filename='data.txt', attachment_id='att1', size=10)]
    gmail = Gmail([message])
    def attachment(*args, **kwargs):
        gmail.account = 'other@example.test'
        return b'not retained'
    gmail.read_attachment = attachment
    review = EmailReview(tmp_path, gmail)
    with pytest.raises(EmailReviewError, match='account changed'):
        review.prepare('after:1')
    assert not list(review.base.glob('review-*'))


def test_pdf_attachment_is_extracted_using_library_and_searchable(tmp_path):
    import pymupdf
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), 'The renewal appointment is 22 October at 14:30.')
    raw = document.tobytes()
    document.close()
    message = email()
    message['attachments'] = [dict(filename='appointment.pdf', size=len(raw), attachment_id='pdf1')]
    gmail = Gmail([message]); gmail.attachment_bytes = raw
    review = EmailReview(tmp_path, gmail)
    result = review.prepare('after:1')
    assert result['attachment_readable'] == 1
    assert '22 October at 14:30' in review.search('renewal appointment')['context']


def test_scanned_pdf_requires_manual_check_instead_of_claiming_it_was_read(tmp_path):
    import pymupdf
    document = pymupdf.open(); document.new_page()
    raw = document.tobytes(); document.close()
    message = email()
    message['attachments'] = [dict(filename='scanned.pdf', size=len(raw), attachment_id='pdf1')]
    gmail = Gmail([message]); gmail.attachment_bytes = raw
    result = EmailReview(tmp_path, gmail).prepare('after:1')
    assert result['messages'][0]['attachments'][0]['status'] == 'no_readable_text'
    assert result['attachment_readable'] == 0
    assert 'open the relevant email' in result['answer']


def test_docx_attachment_includes_paragraphs_and_table_cells(tmp_path):
    import io
    from docx import Document
    document = Document(); document.add_paragraph('Project Merlin contract summary.')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Deposit'; table.cell(0, 1).text = 'GBP 400'
    buffer = io.BytesIO(); document.save(buffer); raw = buffer.getvalue()
    message = email()
    message['attachments'] = [dict(filename='contract.docx', size=len(raw), attachment_id='docx1')]
    gmail = Gmail([message]); gmail.attachment_bytes = raw
    review = EmailReview(tmp_path, gmail)
    result = review.prepare('after:1')
    assert result['attachment_readable'] == 1
    assert 'Deposit | GBP 400' in review.search('contract deposit')['context']


def test_incomplete_participant_headers_cannot_prove_an_exhaustive_fuzzy_search(tmp_path):
    message = email(sender='Someone Else <else@example.test>')
    message['participant_headers_complete'] = False
    result = EmailReview(tmp_path, Gmail([message])).prepare('after:1', participant='Jose Angel')
    assert result['manifest']['matched_total'] == 0
    assert result['manifest']['search_complete'] is False
    assert result['evidence']['complete'] is False
    assert 'matching emails may be missing' in result['answer']
