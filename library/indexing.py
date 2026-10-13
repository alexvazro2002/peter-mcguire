"""Non-destructive text extraction and staged, backwards-compatible FAISS indexing."""
from __future__ import annotations

import json
import pickle
import re
import shutil
import tempfile
import time
from pathlib import Path

from library.paths import LibraryPaths

SUPPORTED = {".txt", ".md", ".pdf", ".docx", ".epub", ".html", ".htm"}
EXCLUDED = {"_KNOWLEDGE_BASE", "ai_models", ".venv", "__pycache__", ".git", ".trash", ".history"}


def source_files(root: Path) -> list[Path]:
    return sorted(p for p in root.rglob("*") if p.is_file() and not p.is_symlink()
                  and p.suffix.lower() in SUPPORTED
                  and not EXCLUDED.intersection(p.relative_to(root).parts))


def make_chunks(text: str, max_chars: int = 1800, overlap: int = 180) -> list[str]:
    """Keep short paragraphs and split long ones; never discard short documents."""
    if max_chars < 50 or not 0 <= overlap < max_chars:
        raise ValueError("Invalid chunk size / overlap.")
    text = re.sub(r"[ \t]+", " ", text.replace("\r\n", "\n").replace("\r", "\n")).strip()
    result, start = [], 0
    while start < len(text):
        end = min(start + max_chars, len(text))
        if end < len(text):
            boundary = max(text.rfind("\n", start + max_chars // 2, end),
                           text.rfind(" ", start + max_chars // 2, end))
            if boundary > start:
                end = boundary
        chunk = text[start:end].strip()
        if chunk:
            result.append(chunk)
        if end >= len(text):
            break
        start = max(start + 1, end - overlap)
    return result


def extract(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8-sig", errors="replace")
    if suffix == ".pdf":
        import pymupdf
        with pymupdf.open(path) as doc:
            return "\n\n".join(f"[PAGE {i + 1}]\n{page.get_text()}" for i, page in enumerate(doc))
    if suffix == ".docx":
        from docx import Document
        doc = Document(path)
        paragraphs = [p.text for p in doc.paragraphs]
        paragraphs += [" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows]
        return "\n\n".join(paragraphs)
    from bs4 import BeautifulSoup
    def html_text(raw):
        soup = BeautifulSoup(raw, "html.parser")
        for tag in soup(["script", "style"]):
            tag.decompose()
        return soup.get_text("\n")
    if suffix == ".epub":
        import ebooklib
        from ebooklib import epub
        book = epub.read_epub(str(path))
        return "\n\n".join(html_text(item.get_content()) for item in book.get_items()
                            if item.get_type() == ebooklib.ITEM_DOCUMENT)
    return html_text(path.read_text(encoding="utf-8-sig", errors="replace"))


def load_metadata(index_dir: Path) -> list[dict]:
    if (index_dir / "metadata.json").exists():
        data = json.loads((index_dir / "metadata.json").read_text(encoding="utf-8"))
    else:
        class PrimitiveOnlyUnpickler(pickle.Unpickler):
            def find_class(self, module, name):
                raise ValueError("Unsafe legacy metadata: object deserialization is disabled.")
        with (index_dir / "metadata.pkl").open("rb") as handle:
            data = PrimitiveOnlyUnpickler(handle).load()
    if not isinstance(data, list) or not data or any(
        not isinstance(item, dict) or not isinstance(item.get("text"), str) or
        not isinstance(item.get("source"), str) or not isinstance(item.get("chunk"), int)
        for item in data
    ):
        raise ValueError("Invalid index metadata. Build a new index from your processed text.")
    return data


def embedding_model(settings: dict, saved: str | None = None):
    from sentence_transformers import SentenceTransformer
    configured = settings.get("embedding_model", "all-MiniLM-L6-v2")
    name = configured if Path(configured).is_dir() else saved or configured
    try:
        return SentenceTransformer(name, local_files_only=settings.get("offline_models", True))
    except OSError as error:
        raise RuntimeError(f"Embedding model '{name}' is not available locally. Set its existing folder "
                           "in Settings, or explicitly turn off Offline models to download it.") from error


def source_manifest(root: Path, files=None):
    """Cheap change detection; content hashing is reserved for loaded search metadata."""
    result = []
    for path in files if files is not None else source_files(root):
        stat = path.stat()
        label = path.with_suffix(".source")
        label_stat = label.stat() if label.is_file() else None
        result.append([str(path.relative_to(root)), stat.st_size, stat.st_mtime_ns,
                       [label_stat.st_size, label_stat.st_mtime_ns] if label_stat else None])
    return result


def index_changed(paths):
    config = paths.index / "config.json"
    if not config.is_file():
        return False
    saved = json.loads(config.read_text(encoding="utf-8"))
    # Legacy indexes remain usable until the user explicitly updates them.
    return "source_manifest" in saved and saved["source_manifest"] != source_manifest(paths.text)


def build_index(paths: LibraryPaths, settings: dict, log=print, progress=None) -> None:
    import faiss
    import numpy as np
    files = source_files(paths.text)
    report = progress or (lambda **event: None)
    if not files:
        raise ValueError("No supported source files found. Select processed .txt files or a document folder.")
    records, failures = [], []
    manifest = source_manifest(paths.text, files)
    report(phase="reading", completed=0, total=len(files), message="Reading documents", passages=0)
    for i, path in enumerate(files, 1):
        relative = path.relative_to(paths.text)
        log(f"Reading {i}/{len(files)}: {relative}")
        report(phase="reading", completed=i - 1, total=len(files), message="Reading documents", current_file=str(relative), passages=len(records))
        try:
            text = extract(path)
            # Scanned PDFs require the user's existing OCR pipeline; do not silently index page labels.
            if not re.sub(r"\[PAGE \d+\]", "", text).strip():
                raise ValueError("No readable text. Run OCR first and select the processed text directory.")
            label = path.with_suffix(".source")
            source = label.read_text(encoding="utf-8").strip() if label.is_file() else str(relative)
            for chunk_id, chunk in enumerate(make_chunks(text)):
                records.append({"id": len(records), "source": source, "chunk": chunk_id, "text": chunk})
        except Exception as error:
            failures.append(str(relative))
            log(f"Skipped {relative}: {error}")
        report(phase="reading", completed=i, total=len(files), message="Reading documents", current_file=str(relative), passages=len(records), skipped=len(failures))
    if not records:
        raise ValueError("No usable text extracted. Existing index has not been changed.")
    report(phase="model", completed=0, total=None, message="Loading your local embedding model", documents=len(files) - len(failures), passages=len(records))
    model = embedding_model(settings)
    log(f"Embedding {len(records):,} passages from {len(files) - len(failures)} documents…")
    vectors = None
    report(phase="embedding", completed=0, total=len(records), message="Embedding document passages", documents=len(files) - len(failures))
    for start in range(0, len(records), 32):
        batch = records[start:start + 32]
        encoded = np.asarray(model.encode([r["text"] for r in batch], batch_size=32,
                                          normalize_embeddings=True, show_progress_bar=False), dtype="float32")
        if encoded.ndim != 2 or encoded.shape[0] != len(batch) or not np.isfinite(encoded).all():
            raise ValueError("Embedding model returned invalid vectors. Existing index was preserved.")
        if vectors is None:
            vectors = np.empty((len(records), encoded.shape[1]), dtype="float32")
        vectors[start:start + len(batch)] = encoded
        report(phase="embedding", completed=start + len(batch), total=len(records), message="Embedding document passages",
               current_file=batch[-1]["source"], documents=len(files) - len(failures))
    report(phase="saving", completed=0, total=None, message="Saving and verifying the new index")
    index = faiss.IndexFlatIP(vectors.shape[1])
    index.add(vectors)
    paths.index.parent.mkdir(parents=True, exist_ok=True)
    stage = Path(tempfile.mkdtemp(prefix=".index-build-", dir=paths.index.parent))
    backup = paths.index.with_name(f"index.backup-{time.time_ns()}")
    try:
        faiss.write_index(index, str(stage / "library.faiss"))
        (stage / "metadata.json").write_text(json.dumps(records, ensure_ascii=False), encoding="utf-8")
        (stage / "config.json").write_text(json.dumps({
            "embedding_model": settings.get("embedding_model", "all-MiniLM-L6-v2"),
            "dimension": int(vectors.shape[1]), "chunks": len(records),
            "documents": len(files) - len(failures), "source": str(paths.source),
            "skipped": failures, "format_version": 3, "source_manifest": manifest,
        }, indent=2), encoding="utf-8")
        if faiss.read_index(str(stage / "library.faiss")).ntotal != len(records):
            raise ValueError("Index verification failed. Existing index was preserved.")
        if paths.index.exists():
            paths.index.rename(backup)
        try:
            stage.rename(paths.index)
        except Exception:
            if backup.exists():
                backup.rename(paths.index)
            raise
    finally:
        if stage.exists():
            shutil.rmtree(stage)  # Only our exact private staging directory.
    log(f"Index ready: {len(records):,} passages. Original files were not modified.")
    report(phase="indexed", completed=len(records), total=len(records), message="Documents indexed", documents=len(files) - len(failures), skipped=len(failures))
